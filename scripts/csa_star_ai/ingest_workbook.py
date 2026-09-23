# Copyright (c) 2026 Guruprasath Annadurai
# SPDX-License-Identifier: MIT

"""Ingest pristine CSA AI-CAIQ v1.1 upstream workbook into remediation_ledger.json."""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

_REPO = Path(__file__).resolve().parents[2]
if str(_REPO) not in sys.path:
    sys.path.insert(0, str(_REPO))

from scripts.csa_star_ai.ledger import ControlAnswer, RemediationLedger, RemediationLedgerRow

_SOURCE = _REPO / "compliance" / "csa-star-ai" / "source"
_UPSTREAM = _SOURCE / "CSA_AI-CAIQ_v1.1_Official_upstream.xlsx"
_INTEGRITY = _SOURCE / "SOURCE_INTEGRITY.json"
_LEDGER_PATH = _REPO / "compliance" / "csa-star-ai" / "ledger" / "remediation_ledger.json"
_MATRIX_PATH = _REPO / "compliance" / "csa-star-ai" / "CONTROL_MATRIX.md"
_EXPECTED_ROWS = 320

_QUESTION_ID_RE = re.compile(r"^[A-Z][A-Za-z0-9&]+-\d+\.\d+$")


def _find_header_row(rows: list[tuple]) -> tuple[int, dict[str, int]]:
    for idx, row in enumerate(rows):
        if not row:
            continue
        headers = [str(c or "").strip().lower() for c in row]
        if "question id" in headers or "question_id" in headers:
            mapping = {h: i for i, h in enumerate(headers)}
            return idx, mapping
    raise ValueError("Could not locate header row with 'Question ID'")


def _col(mapping: dict[str, int], *names: str) -> int | None:
    for name in names:
        if name in mapping:
            return mapping[name]
        for key, idx in mapping.items():
            if name in key:
                return idx
    return None


def _sheet_rows(ws) -> list[tuple]:
    return [tuple(c for c in row) for row in ws.iter_rows(values_only=True)]


def _pick_sheet(wb):
    for name in wb.sheetnames:
        lower = name.lower()
        if "caiq" in lower or "questionnaire" in lower or "ai-caiq" in lower:
            return wb[name]
    return wb[wb.sheetnames[-1]]


def ingest(path: Path) -> RemediationLedger:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = _pick_sheet(wb)
    all_rows = _sheet_rows(ws)
    header_idx, mapping = _find_header_row(all_rows)
    qid_col = _col(mapping, "question id", "question_id")
    q_col = _col(mapping, "question")
    domain_col = _col(mapping, "control domain", "domain")
    title_col = _col(mapping, "control title", "title")
    spec_col = _col(mapping, "control specification", "specification")
    if qid_col is None or q_col is None:
        raise ValueError(f"Missing required columns in mapping: {mapping}")

    ledger_rows: list[RemediationLedgerRow] = []
    for row in all_rows[header_idx + 1 :]:
        if not row or qid_col >= len(row):
            continue
        qid = str(row[qid_col] or "").strip()
        if not _QUESTION_ID_RE.match(qid):
            continue
        question = str(row[q_col] or "").strip()
        domain = ""
        if domain_col is not None and domain_col < len(row):
            domain = str(row[domain_col] or "").strip()
        if not domain and "-" in qid:
            domain = qid.split("-", 1)[0]
        title = str(row[title_col] or "").strip() if title_col is not None and title_col < len(row) else ""
        spec = str(row[spec_col] or "").strip() if spec_col is not None and spec_col < len(row) else ""
        ledger_rows.append(
            RemediationLedgerRow(
                control_id=qid,
                question_id=qid,
                domain=domain,
                control_title=title,
                control_specification=spec,
                question=question,
                response=ControlAnswer.UNASSESSED,
                status="unassessed",
            )
        )

    return RemediationLedger(rows=ledger_rows)


def write_matrix(ledger: RemediationLedger) -> None:
    lines = [
        "# AI-CAIQ Control Matrix (authoritative ingest)",
        "",
        f"Rows: {len(ledger.rows)}",
        f"Summary: {json.dumps(ledger.summary())}",
        "",
        "| Question ID | Domain | Response | Evidence strength | Question (truncated) |",
        "|-------------|--------|----------|-------------------|----------------------|",
    ]
    for row in ledger.rows:
        q = row.question.replace("|", "/").replace("\n", " ")[:100]
        lines.append(
            f"| {row.question_id} | {row.domain} | {row.response.value} | "
            f"{row.evidence_strength.value} | {q} |"
        )
    _MATRIX_PATH.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    path = _UPSTREAM
    if not path.is_file():
        print(f"Missing pristine upstream workbook: {path}", file=sys.stderr)
        integrity = json.loads(_INTEGRITY.read_text(encoding="utf-8"))
        for step in integrity.get("manual_steps", []):
            print(f"  - {step}", file=sys.stderr)
        raise SystemExit(2)

    ledger = ingest(path)
    if len(ledger.rows) != _EXPECTED_ROWS:
        print(
            f"Parser integrity failure: expected {_EXPECTED_ROWS} rows, got {len(ledger.rows)}",
            file=sys.stderr,
        )
        raise SystemExit(1)

    if _INTEGRITY.is_file():
        meta = json.loads(_INTEGRITY.read_text(encoding="utf-8"))
        ledger.upstream_workbook_sha256 = meta.get("workbook_sha256")

    _LEDGER_PATH.parent.mkdir(parents=True, exist_ok=True)
    _LEDGER_PATH.write_text(ledger.model_dump_json(indent=2), encoding="utf-8")
    write_matrix(ledger)
    print(f"Ingested {len(ledger.rows)} controls → {_LEDGER_PATH}")


if __name__ == "__main__":
    main()
