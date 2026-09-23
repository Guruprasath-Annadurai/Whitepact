# Copyright (c) 2026 Guruprasath Annadurai
# SPDX-License-Identifier: MIT

"""Export FINAL AI-CAIQ workbook from remediation ledger (respondent columns only)."""

from __future__ import annotations

import argparse
import json
import shutil
import sys
from pathlib import Path

_REPO = Path(__file__).resolve().parents[2]
_SOURCE = _REPO / "compliance" / "csa-star-ai" / "source"
_UPSTREAM = _SOURCE / "CSA_AI-CAIQ_v1.1_Official_upstream.xlsx"
_LEDGER = _REPO / "compliance" / "csa-star-ai" / "ledger" / "remediation_ledger.json"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    if not _UPSTREAM.is_file():
        print("Pristine upstream workbook required.", file=sys.stderr)
        raise SystemExit(2)
    if not _LEDGER.is_file():
        print("Run ingest_workbook.py and assess_evidence.py first.", file=sys.stderr)
        raise SystemExit(2)
    ledger = json.loads(_LEDGER.read_text(encoding="utf-8"))
    by_id = {r["question_id"]: r for r in ledger.get("rows", [])}
    unassessed = sum(1 for r in ledger.get("rows", []) if r.get("response") == "UNASSESSED")
    if unassessed:
        print(f"Refusing export: {unassessed} rows still UNASSESSED", file=sys.stderr)
        raise SystemExit(1)

    import openpyxl

    shutil.copy2(_UPSTREAM, args.output)
    wb = openpyxl.load_workbook(args.output)
    ws = wb[wb.sheetnames[-1]]
    header_row = None
    for idx, row in enumerate(ws.iter_rows(min_row=1, max_row=20), start=1):
        vals = [str(c.value or "").strip().lower() for c in row]
        if "question id" in vals:
            header_row = idx
            col_map = {v: i for i, v in enumerate(vals)}
            break
    if header_row is None:
        raise SystemExit("Could not find header row")

    ans_idx = col_map.get("csp caiq answer") or col_map.get("caiq answer") or 2
    ssrm_idx = col_map.get("ssrm control ownership") or 3
    impl_idx = col_map.get("csp implementation description (optional/recommended)") or 4
    cust_idx = col_map.get("csc responsibilities  (optional/recommended)") or col_map.get(
        "csc responsibilities (optional/recommended)"
    )

    for row in ws.iter_rows(min_row=header_row + 1):
        qid = str(row[0].value or "").strip()
        if qid not in by_id:
            continue
        entry = by_id[qid]
        resp = entry.get("response", "NO")
        if ans_idx < len(row):
            row[ans_idx].value = resp if resp != "UNASSESSED" else "No"
        if ssrm_idx is not None and ssrm_idx < len(row):
            row[ssrm_idx].value = entry.get("ssrm_owner") or entry.get("implementation_owner") or row[ssrm_idx].value
        if impl_idx is not None and impl_idx < len(row):
            row[impl_idx].value = entry.get("implementation_description") or entry.get("justification") or ""
        if cust_idx is not None and cust_idx < len(row):
            row[cust_idx].value = entry.get("customer_responsibility") or row[cust_idx].value
    wb.save(args.output)
    print(f"Wrote {args.output}")


if __name__ == "__main__":
    main()
